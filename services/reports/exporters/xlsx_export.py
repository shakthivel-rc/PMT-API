"""XLSX export with openpyxl: a Summary sheet (the facts and the KPIs), the report's table,
then one sheet per section. Durations and dates are real numbers and dates with a number
format, so the spreadsheet can sum them; text cells are guarded as in the CSV."""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from services.reports.exporters.formatting import _as_date, footer_lines, guard, tables
from utils import report_math as rm

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXTENSION = "xlsx"

_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E8EAF0")
#: Excel's sheet-name limit, and the characters it refuses in one.
_SHEET_NAME_MAX = 31
_SHEET_FORBIDDEN = set("[]:*?/\\")


def _sheet_name(title: str, used: set[str]) -> str:
    clean = "".join("-" if ch in _SHEET_FORBIDDEN else ch for ch in title).strip() or "Sheet"
    name = clean[:_SHEET_NAME_MAX]
    counter = 2
    while name.lower() in used:
        suffix = f" ({counter})"
        name = clean[:_SHEET_NAME_MAX - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


def _write_cell(sheet, row: int, col: int, value, kind: str, options: dict):
    cell = sheet.cell(row=row, column=col)
    if value is None:
        return cell
    duration_format = options.get("duration_format", "decimal")
    if kind == "duration":
        seconds = float(value)
        if duration_format == "decimal":
            cell.value = round(seconds / 3600, 2)
            cell.number_format = "0.00"
        elif duration_format == "hhmm":
            cell.value = seconds / 86400
            cell.number_format = "[h]:mm"
        elif duration_format == "seconds":
            cell.value = int(seconds)
        else:
            cell.value = rm.format_duration(seconds, "iso")
    elif kind in ("date", "datetime"):
        parsed = _as_date(value)
        if isinstance(parsed, datetime):
            cell.value = parsed.replace(tzinfo=None)
            cell.number_format = "yyyy-mm-dd hh:mm"
        elif isinstance(parsed, date):
            cell.value = parsed
            cell.number_format = "yyyy-mm-dd"
        else:
            cell.value = guard(str(value))
    elif kind == "boolean":
        cell.value = "Yes" if value else "No"
    elif kind in ("number", "percent") and isinstance(value, (int, float)):
        cell.value = value
    else:
        cell.value = guard(str(value))
    return cell


def render(result: dict, options: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    used = {"summary"}
    row = 1
    summary.cell(row=row, column=1, value=guard(str(result["meta"].get("title", "")))).font = (
        Font(bold=True, size=14))
    row += 2
    for label, value in footer_lines(result, options):
        summary.cell(row=row, column=1, value=label).font = _BOLD
        summary.cell(row=row, column=2, value=guard(value))
        row += 1
    if result.get("kpis"):
        row += 1
        summary.cell(row=row, column=1, value="Key figures").font = _BOLD
        row += 1
        for item in result["kpis"]:
            summary.cell(row=row, column=1, value=guard(str(item["label"])))
            _write_cell(summary, row, 2, item.get("value"), item.get("type", "number"), options)
            row += 1
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 48

    for table in tables(result):
        sheet = workbook.create_sheet(_sheet_name(str(table["title"]), used))
        columns = table["columns"]
        for index, col in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=guard(str(col["label"])))
            cell.font = _BOLD
            cell.fill = _HEADER_FILL
            sheet.column_dimensions[get_column_letter(index)].width = max(
                10, min(48, len(str(col["label"])) + 4 if col["type"] != "text" else 36))
        line = 2
        for item in table["rows"]:
            for index, col in enumerate(columns, start=1):
                value = item.get(col["key"])
                if index == 1 and item.get("_depth"):
                    value = "    " * int(item["_depth"]) + str(value or "")
                    sheet.cell(row=line, column=index, value=guard(value))
                    continue
                _write_cell(sheet, line, index, value, col["type"], options)
            line += 1
        totals = table.get("totals")
        if totals:
            for index, col in enumerate(columns, start=1):
                if col["key"] in totals:
                    cell = _write_cell(sheet, line, index, totals[col["key"]], col["type"],
                                       options)
                    cell.font = _BOLD
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
